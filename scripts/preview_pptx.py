"""Geometry preview of a PPTX using real Calibri glyph metrics (a development aid).

This is NOT PowerPoint rendering: it redraws shapes, connectors, the logo and text with Pillow so
layout, wrapping and overflow can be inspected where no Office/LibreOffice renderer exists.
Charts are drawn as simple bars. Usage: python scripts/preview_pptx.py deck.pptx out_dir
"""

from __future__ import annotations

import sys
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

DPI = 150
FONTS = Path("C:/Windows/Fonts")


def _font(size_pt: float, bold: bool) -> ImageFont.FreeTypeFont:
    name = "calibrib.ttf" if bold else "calibri.ttf"
    return ImageFont.truetype(str(FONTS / name), max(1, round(size_pt / 72 * DPI)))


def _rgb(color_format) -> tuple[int, int, int] | None:
    try:
        if color_format.type is None:
            return None
        return tuple(color_format.rgb)  # type: ignore[return-value]
    except (AttributeError, TypeError):
        return None


def _wrap(draw, text: str, font, width_px: float) -> list[str]:
    lines: list[str] = []
    for raw in text.split("\n"):
        current = ""
        for word in raw.split(" "):
            trial = word if not current else current + " " + word
            if draw.textlength(trial, font=font) <= width_px or not current:
                current = trial
            else:
                lines.append(current)
                current = word
        lines.append(current)
    return lines


def render(path: Path, out_dir: Path) -> list[str]:
    prs = Presentation(str(path))
    out_dir.mkdir(parents=True, exist_ok=True)
    problems: list[str] = []
    px = lambda emu: emu / 914400 * DPI  # noqa: E731
    for number, slide in enumerate(prs.slides, 1):
        img = Image.new(
            "RGB", (round(px(prs.slide_width)), round(px(prs.slide_height))), (5, 12, 28)
        )
        draw = ImageDraw.Draw(img)
        for shape in slide.shapes:
            x0, y0 = px(shape.left), px(shape.top)
            x1, y1 = x0 + px(shape.width), y0 + px(shape.height)
            kind = shape.shape_type
            if kind == MSO_SHAPE_TYPE.PICTURE:
                pic = Image.open(__import__("io").BytesIO(shape.image.blob)).convert("RGBA")
                pic = pic.resize((max(1, round(x1 - x0)), max(1, round(y1 - y0))))
                img.paste(pic, (round(x0), round(y0)), pic)
                continue
            if kind == MSO_SHAPE_TYPE.LINE or shape.__class__.__name__ == "Connector":
                color = _rgb(shape.line.color) or (45, 212, 191)
                draw.line(
                    [(x0, y0), (x1, y1)],
                    fill=color,
                    width=max(1, round(shape.line.width / 12700 * DPI / 72)),
                )
                continue
            if shape.has_chart if hasattr(shape, "has_chart") else False:
                chart = shape.chart
                draw.rectangle([x0, y0, x1, y1], outline=(39, 64, 107))
                series = list(chart.plots[0].series)
                cats = list(chart.plots[0].categories)
                top = max((max(s.values) for s in series), default=1) or 1
                gw = (x1 - x0) / max(1, len(cats))
                for ci, cat in enumerate(cats):
                    for si, s in enumerate(series):
                        v = s.values[ci]
                        bw = gw * 0.3
                        bx = x0 + ci * gw + gw * 0.15 + si * bw
                        bh = (y1 - y0 - 60) * v / top
                        col = _rgb(s.format.fill.fore_color) or (200, 200, 200)
                        draw.rectangle([bx, y1 - 40 - bh, bx + bw, y1 - 40], fill=col)
                    draw.text(
                        (x0 + ci * gw + 6, y1 - 30),
                        str(cat),
                        fill=(232, 238, 248),
                        font=_font(10, False),
                    )
                continue
            fill = _rgb(shape.fill.fore_color) if shape.fill.type == 1 else None
            line = _rgb(shape.line.color) if shape.line.fill.type == 1 else None
            try:
                ast = shape.auto_shape_type
            except (ValueError, AttributeError):
                ast = None
            if fill or line:
                if ast is not None and "OVAL" in str(ast):
                    draw.ellipse([x0, y0, x1, y1], fill=fill, outline=line)
                else:
                    r = min(14, (y1 - y0) / 2) if ast is not None and "ROUNDED" in str(ast) else 0
                    draw.rounded_rectangle(
                        [x0, y0, x1, y1], radius=r, fill=fill, outline=line, width=2
                    )
            if not shape.has_text_frame or not shape.text_frame.text.strip():
                continue
            frame = shape.text_frame
            ml, mr = px(frame.margin_left), px(frame.margin_right)
            mt, mb = px(frame.margin_top), px(frame.margin_bottom)
            width = (x1 - x0) - ml - mr
            blocks: list[tuple] = []
            total = 0.0
            for para in frame.paragraphs:
                runs = para.runs
                if not runs:
                    continue
                size = runs[0].font.size.pt if runs[0].font.size else 14
                font = _font(size, bool(runs[0].font.bold))
                color = _rgb(runs[0].font.color) or (232, 238, 248)
                ppr = para._p.pPr
                bullet = (
                    ppr is not None
                    and ppr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}buChar")
                    is not None
                )
                indent = 16 / 72 * DPI if bullet else 0
                text = "".join(r.text for r in runs)
                lines = _wrap(draw, text, font, width - indent)
                after = (para.space_after.pt if para.space_after is not None else 0) / 72 * DPI
                height = len(lines) * size * 1.2 / 72 * DPI
                blocks.append((lines, font, color, height, after, int(indent), para.alignment))
                total += height + after
            avail = (y1 - y0) - mt - mb
            if total > avail * 1.03:
                problems.append(
                    f"slide {number}: '{shape.name}' text {total / DPI:.2f}in "
                    f"> box {avail / DPI:.2f}in"
                )
            anchor = frame.vertical_anchor
            y = y0 + mt
            if anchor is not None and "MIDDLE" in str(anchor):
                y = y0 + mt + max(0.0, (avail - total) / 2)
            for lines, font, color, height, after, indent, align in blocks:
                lh = height / len(lines)
                for i, ln in enumerate(lines):
                    w = draw.textlength(ln, font=font)
                    tx = x0 + ml + indent
                    if "CENTER" in str(align):
                        tx = x0 + ml + (width - w) / 2
                    elif "RIGHT" in str(align):
                        tx = x1 - mr - w
                    draw.text((tx, y), ln, fill=color, font=font)
                    if indent and i == 0:
                        draw.text((x0 + ml, y), "•", fill=color, font=font)
                    y += lh
                y += after
        img.save(out_dir / f"slide_{number}.png")
    return problems


if __name__ == "__main__":
    for line in render(Path(sys.argv[1]), Path(sys.argv[2])):
        sys.stdout.write(line + "\n")
