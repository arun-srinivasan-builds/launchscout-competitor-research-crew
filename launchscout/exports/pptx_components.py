"""Reusable, content-agnostic PowerPoint building blocks (python-pptx, all native shapes).

Nothing here knows about a particular business idea. Layout functions in `pptx_builder` compose
these into slides. Every text element is fitted to its box: the font shrinks to a floor and
then the text is shortened with an ellipsis (the full text always goes to the speaker notes).
"""

from __future__ import annotations

from collections.abc import Sequence
from dataclasses import dataclass, field

from lxml import etree
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

# ---- palette: warm ivory paper, deep navy ink, burgundy brand accent ---------------------------
CREAM = RGBColor(0xFA, 0xF6, 0xEC)  # slide background
PAPER = RGBColor(0xFF, 0xFD, 0xF8)  # cards
INK = RGBColor(0x14, 0x23, 0x3F)  # deep navy text
MUTED = RGBColor(0x5B, 0x67, 0x7A)
BORDER = RGBColor(0xE0, 0xD5, 0xBE)
BURGUNDY = RGBColor(0x7B, 0x1E, 0x3A)
BURGUNDY_LIGHT = RGBColor(0xF3, 0xE3, 0xE7)
TEAL = RGBColor(0x1C, 0x7C, 0x7A)
TEAL_LIGHT = RGBColor(0xDC, 0xEF, 0xEC)
GOLD = RGBColor(0x9C, 0x72, 0x25)
GOLD_LIGHT = RGBColor(0xF5, 0xEC, 0xD3)
BLUE = RGBColor(0x34, 0x69, 0x9A)
BLUE_LIGHT = RGBColor(0xDD, 0xE9, 0xF3)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
TEAL_ON_INK = RGBColor(0x7F, 0xD6, 0xCF)
FONT = "Calibri"
HEAD_FONT = "Georgia"

# ---- geometry (inches) -----------------------------------------------------------------------
SLIDE_W, SLIDE_H = 13.333, 7.5
MARGIN_X = 0.45
CONTENT_W = SLIDE_W - 2 * MARGIN_X
FOOTER_TOP = 7.04
CONTENT_BOTTOM = 6.92
GAP = 0.18

SERIF_EM_FACTOR = 1.12  # Georgia runs wider than Calibri at the same point size
GLYPH_EM = 0.47  # average Calibri glyph width (incl. capitals and bold) as a fraction of the size
LINE_SPACING = 1.2
MIN_SIZE = 11


def estimate_lines(text: str, width_pt: float, size_pt: float) -> int:
    """Wrapped line count estimate (word-aware) shared by layout and the bounds check."""
    per_line = max(1, int(width_pt / (size_pt * GLYPH_EM)))
    lines = 0
    for raw in text.split("\n") or [""]:
        count, current = 1, 0
        for word in raw.split(" "):
            need = len(word) + (1 if current else 0)
            if current + need <= per_line:
                current += need
            else:
                count += 1 if current else 0
                current = len(word)
                while current > per_line:  # an over-long token wraps by character
                    count += 1
                    current -= per_line
        lines += count
    return max(1, lines)


@dataclass
class Para:
    text: str
    size: float = 13
    color: RGBColor = INK
    bold: bool = False
    bullet: bool = False
    after: float = 4  # points of space after
    align: PP_ALIGN = PP_ALIGN.LEFT
    italic: bool = False
    font: str = FONT


@dataclass
class Fit:
    paras: list[Para] = field(default_factory=list)
    truncated: bool = False


BULLET_INDENT_PT = 16.0


def _needed_pt(paras: Sequence[Para], width_pt: float) -> float:
    total = 0.0
    for p in paras:
        indent = BULLET_INDENT_PT if p.bullet else 0.0
        eff = p.size * (SERIF_EM_FACTOR if p.font == HEAD_FONT else 1.0)
        lines = estimate_lines(p.text, width_pt - indent, eff)
        total += lines * p.size * LINE_SPACING + p.after
    return total


_BOUNDARIES = (". ", "; ", " — ", ": ", ", ")


def _shorten(text: str, max_chars: int) -> str:
    """Shorten to a complete clause where possible; an ellipsis is the last resort."""
    if len(text) <= max_chars:
        return text
    window = text[: max(1, max_chars)]
    floor = int(max_chars * 0.4)
    best = -1
    for sep in _BOUNDARIES:
        at = window.rfind(sep)
        if at >= floor:
            best = max(best, at)
    if best >= 0:
        cut = window[:best].rstrip(",;:.- ")
        if cut:
            return cut + ("." if text[best : best + 2] == ". " else "")
    cut = text[: max(1, max_chars - 1)].rsplit(" ", 1)[0].rstrip(",;:.- ")
    return (cut or text[: max_chars - 1]) + "…"


def fit_paragraphs(paras: Sequence[Para], width_pt: float, height_pt: float) -> Fit:
    """Shrink fonts (down to MIN_SIZE) then shorten text until the paragraphs fit the box."""
    work = [Para(**p.__dict__) for p in paras]
    while _needed_pt(work, width_pt) > height_pt and any(p.size > MIN_SIZE for p in work):
        for p in work:  # shrink every paragraph one point, never below the readability floor
            p.size = max(MIN_SIZE, p.size - 1)
    truncated = False
    guard = 0
    while _needed_pt(work, width_pt) > height_pt and guard < 200:
        guard += 1
        longest = max(work, key=lambda p: len(p.text), default=None)
        if longest is None or len(longest.text) <= 30:
            break
        longest.text = _shorten(longest.text, int(len(longest.text) * 0.85))
        truncated = True
    return Fit(work, truncated)


def _apply_run(run, p: Para) -> None:
    run.font.name = p.font
    run.font.size = Pt(p.size)
    run.font.bold = p.bold
    run.font.italic = p.italic
    run.font.color.rgb = p.color


def _set_bullet(paragraph) -> None:
    props = paragraph._p.get_or_add_pPr()
    props.set("marL", str(int(Pt(BULLET_INDENT_PT))))
    props.set("indent", str(-int(Pt(BULLET_INDENT_PT))))
    for tag in ("a:buNone", "a:buChar"):
        for child in props.findall(qn(tag)):
            props.remove(child)
    props.append(props.makeelement(qn("a:buChar"), {"char": "•"}))


def box(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    fill: RGBColor | None = None,
    line: RGBColor | None = None,
    line_w: float = 1.0,
    round_: bool = True,
    shape: MSO_SHAPE | None = None,
    name: str | None = None,
):
    """A filled rectangle / rounded rectangle / other autoshape without text."""
    kind = shape or (MSO_SHAPE.ROUNDED_RECTANGLE if round_ else MSO_SHAPE.RECTANGLE)
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = min(0.5, 0.08 / max(min(w, h), 0.2))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(line_w)
    s.shadow.inherit = False
    if name:
        s.name = name
    return s


def text(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    paras: Sequence[Para] | Para,
    *,
    fill: RGBColor | None = None,
    line: RGBColor | None = None,
    anchor: MSO_ANCHOR = MSO_ANCHOR.TOP,
    pad: float = 0.08,
    round_: bool = True,
    shape: MSO_SHAPE | None = None,
    name: str | None = None,
    line_w: float = 1.0,
    wrap: bool = True,
):
    """A text-bearing shape whose text is fitted to the box. Returns (shape, Fit)."""
    items = [paras] if isinstance(paras, Para) else list(paras)
    if fill is None and line is None and shape is None:
        s = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    else:
        s = box(slide, x, y, w, h, fill=fill, line=line, round_=round_, shape=shape, line_w=line_w)
    frame = s.text_frame
    frame.word_wrap = wrap
    frame.vertical_anchor = anchor
    frame.margin_left = frame.margin_right = Inches(pad)
    frame.margin_top = frame.margin_bottom = Inches(min(pad, 0.06))
    width_pt = (w - 2 * pad) * 72
    height_pt = (h - 2 * min(pad, 0.06)) * 72
    fit = fit_paragraphs(items, width_pt, height_pt)
    for i, p in enumerate(fit.paras):
        para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        para.alignment = p.align
        para.space_after = Pt(p.after)
        if p.bullet:
            _set_bullet(para)
        run = para.add_run()
        run.text = p.text
        _apply_run(run, p)
    if name:
        s.name = name
    return s, fit


def needed_height_in(paras: Sequence[Para], w: float, pad: float = 0.08) -> float:
    """Height (inches) a text box of width `w` needs for these paragraphs."""
    return (_needed_pt(paras, (w - 2 * pad) * 72) / 72) + 2 * min(pad, 0.06) + 0.02


CAPS_EM = 0.65  # bold capitals are wider than average text; LibreOffice also renders wider


def chip_width(label: str, size: float = 11, pad: float = 0.16) -> float:
    """Width that holds the whole label on one line, with padding either side."""
    return len(label) * size * CAPS_EM / 72 + 2 * pad + 0.12


def chip(
    slide,
    x: float,
    y: float,
    label: str,
    fill: RGBColor,
    color: RGBColor,
    *,
    size: float = 11,
    h: float = 0.32,
    name: str | None = None,
    line: RGBColor | None = None,
):
    """A pill-shaped status chip. Returns (shape, width)."""
    w = chip_width(label, size)
    s, _ = text(
        slide,
        x,
        y,
        w,
        h,
        Para(label, size, color, bold=True, align=PP_ALIGN.CENTER, after=0),
        fill=fill,
        line=line,
        anchor=MSO_ANCHOR.MIDDLE,
        pad=0.16,
        shape=MSO_SHAPE.ROUNDED_RECTANGLE,
        name=name,
        wrap=False,  # a status label must never break onto a second line
    )
    s.adjustments[0] = 0.5
    return s, w


def badge(
    slide,
    x: float,
    y: float,
    d: float,
    glyph: str,
    fill: RGBColor,
    color: RGBColor = WHITE,
    *,
    size: float = 12,
    name: str | None = None,
):
    """A round icon badge with a short glyph."""
    s, _ = text(
        slide,
        x,
        y,
        d,
        d,
        Para(glyph, size, color, bold=True, align=PP_ALIGN.CENTER, after=0),
        fill=fill,
        anchor=MSO_ANCHOR.MIDDLE,
        pad=0.0,
        wrap=False,
        shape=MSO_SHAPE.OVAL,
        name=name,
    )
    return s


def arrow(
    slide,
    x1: float,
    y1: float,
    x2: float,
    y2: float,
    color: RGBColor = TEAL,
    *,
    width_pt: float = 2.25,
    head: bool = True,
    name: str | None = None,
):
    """A native connector with an optional arrowhead."""
    c = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    c.line.color.rgb = color
    c.line.width = Pt(width_pt)
    if head:
        ln = c.line._get_or_add_ln()
        tail = etree.SubElement(ln, qn("a:tailEnd"))
        tail.set("type", "triangle")
        tail.set("w", "med")
        tail.set("len", "med")
    if name:
        c.name = name
    return c


def panel(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    heading: str | None = None,
    *,
    accent: RGBColor = BURGUNDY,
    fill: RGBColor = PAPER,
    name: str | None = None,
    reserve: float = 0.0,
) -> float:
    """A bordered card with an accent strip and optional heading. Returns the inner top y."""
    box(slide, x, y, w, h, fill=fill, line=BORDER, name=name)
    box(slide, x, y + 0.12, 0.06, h - 0.24, fill=accent, round_=False, name="Accent")
    inner = y + 0.1
    if heading:
        text(
            slide,
            x + 0.18,
            y + 0.08,
            w - 0.3 - reserve,
            0.32,
            Para(heading.upper(), 11.5, accent, bold=True, after=0),
            pad=0.02,
            name="Panel heading",
        )
        inner = y + 0.44
    return inner


def bar(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    fraction: float,
    color: RGBColor,
    *,
    name: str,
) -> None:
    """A native progress bar: a pale track and a filled share (0..1)."""
    box(slide, x, y, w, h, fill=BORDER, name=f"{name} track")
    share = max(0.0, min(1.0, fraction))
    if share > 0:
        box(slide, x, y, max(w * share, 0.06), h, fill=color, name=f"{name} fill")


def split_widths(total: float, n: int, gap: float = GAP) -> list[float]:
    n = max(1, n)
    each = (total - gap * (n - 1)) / n
    return [each] * n


# ---- pictograms built from native autoshapes (editable; no image assets) -----------------------
def icon(slide, x: float, y: float, d: float, kind: str, fill: RGBColor, *, name: str):
    """A round badge holding a simple white pictogram. Returns the circle."""
    circle = box(slide, x, y, d, d, fill=fill, shape=MSO_SHAPE.OVAL, name=f"{name} icon")
    g = d * 0.46
    gx, gy = x + (d - g) / 2, y + (d - g) / 2

    def glyph(shape, gx_, gy_, w, h, label):
        return box(slide, gx_, gy_, w, h, fill=WHITE, shape=shape, name=f"{name} {label}")

    if kind == "person":
        glyph(MSO_SHAPE.OVAL, gx + g * 0.28, gy, g * 0.44, g * 0.44, "head")
        glyph(MSO_SHAPE.ROUND_2_SAME_RECTANGLE, gx + g * 0.08, gy + g * 0.52, g * 0.84, g * 0.48,
              "body")
    elif kind == "target":
        glyph(MSO_SHAPE.DONUT, gx, gy, g, g, "ring")
        glyph(MSO_SHAPE.OVAL, gx + g * 0.34, gy + g * 0.34, g * 0.32, g * 0.32, "dot")
    elif kind == "chart":
        for i, h in enumerate((0.45, 0.75, 1.0)):
            glyph(MSO_SHAPE.RECTANGLE, gx + i * g * 0.36, gy + g * (1 - h), g * 0.26, g * h,
                  f"bar{i}")
    elif kind == "speech":
        glyph(MSO_SHAPE.ROUNDED_RECTANGULAR_CALLOUT, gx, gy, g, g * 0.82, "bubble")
    elif kind == "doc":
        glyph(MSO_SHAPE.FOLDED_CORNER, gx + g * 0.1, gy, g * 0.8, g, "page")
    elif kind == "bolt":
        glyph(MSO_SHAPE.LIGHTNING_BOLT, gx, gy, g, g, "bolt")
    elif kind == "heart":
        glyph(MSO_SHAPE.HEART, gx, gy + g * 0.05, g, g * 0.9, "heart")
    elif kind == "star":
        glyph(MSO_SHAPE.STAR_5_POINT, gx, gy, g, g, "star")
    elif kind == "flag":
        glyph(MSO_SHAPE.RECTANGLE, gx + g * 0.05, gy, g * 0.1, g, "pole")
        glyph(MSO_SHAPE.RIGHT_TRIANGLE, gx + g * 0.15, gy, g * 0.8, g * 0.55, "flag")
    elif kind == "tag":
        glyph(MSO_SHAPE.PENTAGON, gx, gy + g * 0.15, g, g * 0.7, "tag")
    elif kind == "check":
        glyph(MSO_SHAPE.MATH_MULTIPLY, gx, gy, g, g, "mark")
    else:
        glyph(MSO_SHAPE.OVAL, gx + g * 0.2, gy + g * 0.2, g * 0.6, g * 0.6, "dot")
    return circle
