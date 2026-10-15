"""M6 Pitch Deck: source loading, claim check, labels, PPTX and UI. Offline, no model calls."""

from __future__ import annotations

import json
from datetime import date
from pathlib import Path

import pytest
from pptx import Presentation
from streamlit.testing.v1 import AppTest

from launchscout import deck
from launchscout.runner import set_manager
from launchscout.schemas import (
    Competitor,
    CompetitorSet,
    DeckSpec,
    Price,
    ResearchInput,
    ReviewFinding,
    ReviewSet,
    SlideSpec,
    StrategyOutput,
)
from launchscout.validation import ValidationReport
from tests.conftest import make_manager
from tests.helpers import FAKE_SETTINGS

APP = str(Path(__file__).resolve().parent.parent / "app.py")
U1, U2, U3 = "https://a.example/menu", "https://b.example/", "https://a.example/reviews"
RUN_ID = "20261014-120000-juice-pune-ab12"
REPORT = """# Report

## 1. Idea and Market
Cold-pressed juice for gym-goers in Pune.

## 2. Key Competitors
| Competitor | Price |
|---|---|
| Alpha Juice | ₹120 per bottle (verified) |
| Beta Press | unknown |

## 6. 30-Day Launch Plan
Week 1: sign up 20 gym members. Week 2: pilot. Week 3: refine. Week 4: launch.
"""


def _write_run(runs: Path, state: str = "partial_failed_acceptance", report: bool = True) -> Path:
    root = runs / RUN_ID
    root.mkdir(parents=True)
    ri = ResearchInput(idea="Cold-pressed juice", city="Pune", target_customer="gym-goers")
    (root / "input.json").write_text(ri.model_dump_json(), encoding="utf-8")
    comps = CompetitorSet(
        competitors=[
            Competitor(
                name="Alpha Juice",
                url=U1,
                offer="Daily bottles",
                price=Price(
                    status="verified",
                    amount=120,
                    currency="INR",
                    unit="bottle",
                    evidence_url=U1,
                    accessed_on=date(2026, 10, 14),
                ),
                evidence_urls=[U1],
            ),
            Competitor(name="Beta Press", url=U2, offer="Weekly box", evidence_urls=[U2]),
        ],
        shortfall_reason="only two found",
    )
    (root / "competitors.json").write_text(comps.model_dump_json(), encoding="utf-8")
    reviews = ReviewSet(
        findings=[
            ReviewFinding(
                competitor="Alpha Juice",
                kind="praise",
                text="Fresh taste",
                basis="quoted",
                source_url=U3,
            )
        ]
    )
    (root / "reviews.json").write_text(reviews.model_dump_json(), encoding="utf-8")
    strategy = StrategyOutput(
        market_gap="No gym-focused subscription",
        recommended_offer="Gym-gate pickup plan",
        positioning_line="Fuel after the lift",
        recommended_price="INR 99 per bottle",
        rationale="Both competitors skip gyms",
    )
    (root / "strategy.json").write_text(strategy.model_dump_json(), encoding="utf-8")
    evidence = [
        {"url": u, "kind": "scraped", "accessed_on": "2026-10-14", "text": "x"}
        for u in (U1, U2, U3)
    ]
    (root / "evidence.json").write_text(json.dumps(evidence), encoding="utf-8")
    validation = ValidationReport(
        state=state,
        acceptance={"result": "fail", "reasons": ["only 2 competitors"]},
        verified_prices=1,
    )
    (root / "validation.json").write_text(validation.model_dump_json(), encoding="utf-8")
    metadata = {
        "run_id": RUN_ID,
        "mode": "sequential",
        "started_at": "2026-10-14T12:00:00+00:00",
        "finished_at": "2026-10-14T12:02:00+00:00",
        "status": "complete",
        "acceptance": {"result": "fail", "reasons": []},
        "usage": {},
    }
    (root / "metadata.json").write_text(json.dumps(metadata), encoding="utf-8")
    if report:
        name = "report.md" if state == "validated_pass" else "report-PARTIAL.md"
        (root / name).write_text(REPORT, encoding="utf-8")
    return root


def _good_slides() -> list[dict]:
    return [
        {
            "title": "Idea & Market",
            "bullets": ["Cold-pressed juice for gym-goers in Pune"],
            "notes": "From the report.",
            "source_refs": ["Report: Idea and Market"],
        },
        {
            "title": "Competitor Landscape",
            "bullets": [
                "Alpha Juice: daily bottles at ₹120 per bottle (verified)",
                "Beta Press: weekly box, price unknown",
            ],
            "notes": "",
            "source_refs": [U1, U2],
        },
        {
            "title": "Customer Pain Points",
            "bullets": ["Alpha Juice: fresh taste praised"],
            "notes": "",
            "source_refs": [U3],
        },
        {
            "title": "Market Gap",
            "bullets": ["No gym-focused subscription"],
            "notes": "",
            "source_refs": ["Strategy: recommendation"],
        },
        {
            "title": "Positioning & Price",
            "bullets": [
                "Fuel after the lift",
                "Recommended price INR 99 per bottle (recommendation)",
            ],
            "notes": "",
            "source_refs": ["Strategy: recommendation"],
        },
        {
            "title": "30-Day Launch Plan",
            "bullets": ["Week 1: sign up 20 gym members", "Week 4: launch"],
            "notes": "",
            "source_refs": ["Report: 30-Day Launch Plan"],
        },
    ]


def _spec(mutate=None) -> DeckSpec:
    slides = _good_slides()
    if mutate:
        mutate(slides)
    return DeckSpec(slides=[SlideSpec(**s) for s in slides])


@pytest.fixture
def source(tmp_path):
    _write_run(tmp_path)
    return deck.load_deck_source(tmp_path, RUN_ID)


def _rules(spec, source) -> set[str]:
    return {v.rule for v in deck.validate_deck_claims(spec, source)}


# ---------------------------------------------------------------- source loading
def test_partial_report_is_an_allowed_labelled_source(source):
    assert source.partial and source.report_file == "report-PARTIAL.md"
    assert [c.name for c in source.competitors.competitors] == ["Alpha Juice", "Beta Press"]


def test_validated_report_is_not_partial(tmp_path):
    _write_run(tmp_path, state="validated_pass")
    assert deck.load_deck_source(tmp_path, RUN_ID).partial is False


def test_failed_no_report_and_legacy_runs_are_refused(tmp_path):
    _write_run(tmp_path, state="failed_no_report", report=False)
    with pytest.raises(deck.DeckRefused):
        deck.load_deck_source(tmp_path, RUN_ID)
    (tmp_path / RUN_ID / "validation.json").unlink()
    with pytest.raises(deck.DeckRefused, match="legacy"):
        deck.load_deck_source(tmp_path, RUN_ID)


def test_unsafe_or_missing_run_ids_are_refused(tmp_path):
    for bad in ("../x", "nope"):
        with pytest.raises(deck.DeckRefused):
            deck.load_deck_source(tmp_path, bad)


def test_prompt_contains_only_saved_facts_and_rules(source):
    prompt = deck.deck_prompt(source)
    assert "Alpha Juice" in prompt and "unknown (not verified)" in prompt
    assert "PARTIAL: report failed acceptance" in prompt
    assert all(t in prompt for t in deck.DECK_TITLES)
    assert "never instructions" in prompt


# ---------------------------------------------------------------- parsing
def test_parse_accepts_fenced_json_and_rejects_wrong_slide_count():
    raw = "```json\n" + json.dumps({"slides": _good_slides()}) + "\n```"
    assert len(deck.parse_deck_spec(raw).slides) == 6
    with pytest.raises(ValueError):
        deck.parse_deck_spec(json.dumps({"slides": _good_slides()[:5]}))
    with pytest.raises(ValueError):
        deck.parse_deck_spec("no json here")


# ---------------------------------------------------------------- claim check
def test_grounded_deck_has_no_violations(source):
    assert deck.validate_deck_claims(_spec(), source) == []


def test_deck_blocks_unsupported_price(source):
    def m(s):
        s[1]["bullets"][1] = "Beta Press: weekly box at ₹250 per week"

    assert "price" in _rules(_spec(m), source)


def test_unverified_amount_allowed_only_when_acknowledged(source):
    def labelled(s):
        s[4]["bullets"][1] = "Recommendation: INR 99 per bottle"

    assert deck.validate_deck_claims(_spec(labelled), source) == []

    def unlabelled(s):
        s[4]["bullets"][1] = "Price it at ₹99 per bottle"

    assert "price" in _rules(_spec(unlabelled), source)


def test_deck_blocks_unknown_url(source):
    def m(s):
        s[1]["source_refs"] = ["https://invented.example/page"]

    assert "url" in _rules(_spec(m), source)

    def inline(s):
        s[2]["notes"] = "see https://other.example/x"

    assert "url" in _rules(_spec(inline), source)


def test_deck_blocks_unsourced_statistic(source):
    def m(s):
        s[3]["bullets"].append("Gym-goers are 62% of the market")

    assert "statistic" in _rules(_spec(m), source)


def test_deck_blocks_invented_competitor_and_year(source):
    def m(s):
        s[1]["bullets"].append("Gamma Squeeze: premium bottles")
        s[0]["bullets"].append("Market founded in 1987")

    assert {"competitor", "date"} <= _rules(_spec(m), source)


def test_deck_allows_slide_numbers_dates_metadata(source):
    def m(s):
        s[5]["bullets"] = ["Week 1: sign up 20 gym members", "Week 2: pilot", "Week 4: launch"]

    assert deck.validate_deck_claims(_spec(m), source) == []


def test_wrong_title_is_a_violation(source):
    def m(s):
        s[0]["title"] = "Overview"

    assert "title" in _rules(_spec(m), source)


# ---------------------------------------------------------------- labels
def test_labels_mark_partial_unknown_recommendation_and_assumption(source):
    labelled = deck.apply_labels(_spec(), source)
    assert all("PARTIAL EVIDENCE" in s.labels[0] for s in labelled.slides)
    assert any("Beta Press" in label for label in labelled.slides[1].labels)
    assert "Alpha Juice" not in " ".join(labelled.slides[1].labels)  # verified price: no label
    assert any(
        "Recommendation (not a verified market price)" in x for x in labelled.slides[4].labels
    )
    assert any("Assumption" in x for x in labelled.slides[3].labels)
    assert any("Recommended actions" in x for x in labelled.slides[5].labels)


def test_validated_source_gets_no_partial_label(tmp_path):
    _write_run(tmp_path, state="validated_pass")
    src = deck.load_deck_source(tmp_path, RUN_ID)
    labelled = deck.apply_labels(_spec(), src)
    assert not any("PARTIAL" in label for s in labelled.slides for label in s.labels)


# ---------------------------------------------------------------- generation + export
def _kickoff(slides=None):
    calls: list[str] = []

    def run(prompt: str):
        calls.append(prompt)
        return json.dumps({"slides": slides or _good_slides()}), {
            "basis": "measured",
            "total_tokens": 1234,
        }

    run.calls = calls  # type: ignore[attr-defined]
    return run


def test_generate_uses_one_agent_call_and_exports_editable_pptx(tmp_path, source):
    from launchscout.exports.pptx_builder import SLIDE_NAMES

    kick = _kickoff()
    result = deck.generate_deck(source, kickoff=kick)
    assert len(kick.calls) == 1 and result.violations == []
    path = deck.save_deck(tmp_path, source, result, "test-model", None)
    assert path is not None and path.name == "launchscout-pitch-deck-PARTIAL.pptx"
    prs = Presentation(str(path))
    assert len(prs.slides) == 6
    for slide, title in zip(prs.slides, deck.DECK_TITLES, strict=True):
        pictures = [shape for shape in slide.shapes if shape.shape_type == 13]
        assert [p.name for p in pictures] == ["LS logo"]  # the logo is the only picture
        names = {s.name for s in slide.shapes}
        assert {"Title", "Subtitle", "Labels", "Footer", "Evidence status"} <= names
        shown = next(s for s in slide.shapes if s.name == "Title").text_frame.text
        assert shown in SLIDE_NAMES  # the V3 slide name; the validated deck title is the subtitle
        notes = slide.notes_slide.notes_text_frame.text
        assert "Validated deck title: " + title in notes
        assert slide.notes_slide.notes_text_frame.text is not None
    matrix = " ".join(s.text_frame.text for s in prs.slides[1].shapes if s.has_text_frame)
    assert "Alpha Juice" in matrix and "Beta Press" in matrix
    assert (
        "PARTIAL EVIDENCE"
        in next(s for s in prs.slides[0].shapes if s.name == "Labels").text_frame.text
    )
    assert "Sources:" in prs.slides[1].notes_slide.notes_text_frame.text


def test_violations_block_export_but_keep_spec_and_meta(tmp_path, source):
    bad = _good_slides()
    bad[1]["bullets"].append("Beta Press: now ₹999 per month")
    result = deck.generate_deck(source, kickoff=_kickoff(bad))
    assert result.violations
    assert deck.save_deck(tmp_path, source, result, "m", None) is None
    assert not list(deck.deck_dir(tmp_path, RUN_ID).glob("*.pptx"))
    spec, meta, pptx = deck.load_saved_deck(tmp_path, RUN_ID)
    assert pptx is None and meta["exported"] is False and meta["violations"]


def test_generation_does_not_touch_research_files(tmp_path, source):
    root = tmp_path / RUN_ID
    before = {p.name: p.read_bytes() for p in root.iterdir() if p.is_file()}
    deck.save_deck(tmp_path, source, deck.generate_deck(source, kickoff=_kickoff()), "m", None)
    after = {p.name: p.read_bytes() for p in root.iterdir() if p.is_file()}
    assert before == after


def test_agent_has_no_tools_and_no_delegation(monkeypatch, source):
    captured = {}

    class FakeAgent:
        def __init__(self, **kwargs):
            captured.update(kwargs)

    class FakeCrew:
        usage_metrics = None

        def __init__(self, **kwargs):
            pass

        def kickoff(self):
            return json.dumps({"slides": _good_slides()})

    monkeypatch.setattr("crewai.Agent", FakeAgent)
    monkeypatch.setattr("crewai.Crew", FakeCrew)
    monkeypatch.setattr("crewai.Task", lambda **kw: object())
    monkeypatch.setattr("launchscout.llm.build_llm", lambda s: object())
    result = deck.generate_deck(source, FAKE_SETTINGS)
    assert captured["tools"] == [] and captured["allow_delegation"] is False
    assert captured["role"] == deck.ROLE_DECK
    assert result.usage == {"basis": "unavailable"}


# ---------------------------------------------------------------- UI
def _app(tmp_path) -> AppTest:
    set_manager(make_manager(tmp_path))
    return AppTest.from_file(APP, default_timeout=30).run()


def test_tab_is_disabled_without_saved_runs(tmp_path):
    at = _app(tmp_path)
    buttons = at.tabs[1].button
    assert len(buttons) == 1 and buttons[0].disabled


def test_tab_uses_saved_run_without_model_calls_until_clicked(tmp_path, monkeypatch):
    _write_run(tmp_path / "runs")
    calls: list[str] = []

    def fake_generate(source, settings=None, kickoff=None):
        calls.append(source.run_id)
        return deck.DeckResult(
            spec=deck.apply_labels(_spec(), source), violations=[], usage={"total_tokens": 10}
        )

    monkeypatch.setattr(deck, "generate_deck", fake_generate)
    at = _app(tmp_path)
    assert calls == []  # opening the tab and selecting a run never generates
    tab = at.tabs[1]
    assert tab.selectbox(key="deck_source_run").value == RUN_ID
    assert tab.button(key="deck_generate").disabled is False
    text = " ".join(m.value for m in tab.markdown)
    assert "failed acceptance" in text and "one model call" in text
    tab.button(key="deck_generate").click()
    at.run()
    assert calls == [RUN_ID]
    out = " ".join(m.value for m in at.tabs[1].markdown)
    assert "Editable PPTX ready" in out and "Partial evidence deck" in out
    assert "Competitor Landscape" in out and "PARTIAL EVIDENCE" in out
    again = at.run()  # a rerun shows the stored deck and makes no new call
    assert calls == [RUN_ID] and not again.exception
    assert (tmp_path / "runs" / RUN_ID / "deck" / "launchscout-pitch-deck-PARTIAL.pptx").is_file()


def test_tab_shows_blocked_export_with_reasons(tmp_path, monkeypatch):
    _write_run(tmp_path / "runs")
    bad = _good_slides()
    bad[1]["bullets"].append("Beta Press: now ₹999 per month")
    monkeypatch.setattr(
        deck, "generate_deck", lambda source, settings=None, kickoff=None: _blocked(source, bad)
    )
    at = _app(tmp_path)
    at.tabs[1].button(key="deck_generate").click()
    at.run()
    out = " ".join(m.value for m in at.tabs[1].markdown)
    assert "Claim check blocked the export" in out and "Export blocked" in out


def _blocked(source, slides):
    spec = deck.apply_labels(_spec_from(slides), source)
    return deck.DeckResult(spec=spec, violations=deck.validate_deck_claims(spec, source))


def _spec_from(slides) -> DeckSpec:
    return DeckSpec(slides=[SlideSpec(**s) for s in slides])


# ---------------------------------------------------------------- recommendation wording (fix)
def _launch(bullet: str):
    def m(s):
        s[5]["bullets"] = [bullet]

    return _spec(m)


@pytest.mark.parametrize("word", ["proposed", "suggested", "recommended", "Recommendation:"])
def test_recommendation_wording_accepts_only_the_researched_amount(source, word):
    # strategy.recommended_price is "INR 99 per bottle"
    assert (
        deck.validate_deck_claims(_launch(f"Week 1: validate the {word} INR 99 price"), source)
        == []
    )


def test_recommendation_wording_does_not_excuse_other_amounts(source):
    for text in (
        "Week 1: validate the proposed ₹250 price",  # not the researched recommendation
        "Week 1: validate the suggested ₹120 price",  # a competitor's verified price, wrong class
        "Week 1: validate the price of ₹99",  # right amount, no recommendation wording
    ):
        rules = {v.rule for v in deck.validate_deck_claims(_launch(text), source)}
        assert "price" in rules, text


def test_recommendation_wording_far_from_the_amount_does_not_count(source):
    text = "Recommended next step is a pilot. " + "Detail. " * 20 + "Priced at ₹99 per bottle."
    assert "price" in {v.rule for v in deck.validate_deck_claims(_launch(text), source)}


def test_recommendation_wording_never_excuses_competitor_slides(source):
    def m(s):
        s[1]["bullets"][1] = "Beta Press: proposed INR 99 weekly box"

    assert "price" in _rules(_spec(m), source)


def test_competitor_unverified_price_still_blocked_with_recommendation_word(source):
    def m(s):
        s[1]["bullets"][1] = "Beta Press: recommended ₹250 weekly box"

    assert "price" in _rules(_spec(m), source)


# ---------------------------------------------------------------- offline revalidation + bounds
def test_revalidate_saved_deck_exports_without_model_call_or_research_change(tmp_path, source):
    bad = _good_slides()
    bad[5]["bullets"] = ["Week 1: validate the proposed INR 99 price"]
    bad[1]["bullets"].append("Beta Press: now ₹999 per month")
    result = deck.generate_deck(source, kickoff=_kickoff(bad))
    assert deck.save_deck(tmp_path, source, result, "m", None) is None
    spec, meta, _ = deck.load_saved_deck(tmp_path, RUN_ID)
    original_generated_at = meta["generated_at"]
    # still blocked: the unsupported competitor price remains
    assert deck.revalidate_saved_deck(tmp_path, RUN_ID)
    # fix only the stored text with the unsupported claim removed, then revalidate
    path = deck.deck_dir(tmp_path, RUN_ID) / deck.SPEC_FILE
    data = json.loads(path.read_text(encoding="utf-8"))
    data["slides"][1]["bullets"].pop()
    path.write_text(json.dumps(data), encoding="utf-8")
    before = {p.name: p.read_bytes() for p in (tmp_path / RUN_ID).iterdir() if p.is_file()}
    assert deck.revalidate_saved_deck(tmp_path, RUN_ID) == []
    assert before == {p.name: p.read_bytes() for p in (tmp_path / RUN_ID).iterdir() if p.is_file()}
    _, meta, pptx = deck.load_saved_deck(tmp_path, RUN_ID)
    assert pptx is not None and meta["exported"] is True
    assert meta["generated_at"] == original_generated_at  # original generation record kept
    assert len(Presentation(str(pptx)).slides) == 6


def test_text_bounds_check_passes_normal_deck_and_flags_overflow(tmp_path, source):
    from launchscout.exports.pptx_builder import build_pptx, check_text_bounds

    ok = tmp_path / "ok.pptx"
    build_pptx(deck.apply_labels(_spec(), source), source, ok, date(2026, 10, 14))
    assert check_text_bounds(ok) == []

    prs = Presentation(str(ok))
    victim = next(sh for sh in prs.slides[0].shapes if sh.name == "Business idea")
    victim.height = 20000  # far too small for its text
    broken = tmp_path / "broken.pptx"
    prs.save(str(broken))
    assert any("overflow" in p for p in check_text_bounds(broken))


# ---------------------------------------------------------------- multiline warning banners
LONG_LABEL = (
    "Recommendation (not a verified market price): INR 699 for a seven-day plan of one juice "
    "per day (about INR 100 per juice); proposed price only, with serving size and delivery "
    "terms to be specified before sale and confirmed with a pilot group of consenting customers."
)


def _banner_deck(source):
    labelled = deck.apply_labels(_spec(), source)
    slides = []
    for s in labelled.slides:
        slides.append(s.model_copy(update={"labels": [*s.labels, LONG_LABEL]}))
    return DeckSpec(slides=slides)


def test_multiline_warning_banner_grows_to_contain_its_text(tmp_path, source):
    from pptx.util import Inches

    from launchscout.exports.pptx_builder import build_pptx, check_text_bounds

    path = tmp_path / "banner.pptx"
    build_pptx(_banner_deck(source), source, path, date(2026, 10, 14))
    prs = Presentation(str(path))
    for slide in prs.slides:
        banner = next(s for s in slide.shapes if s.name == "Labels")
        assert banner.height > Inches(0.6)  # wrapped to several lines, not a one-line strip
        chrome = {
            "LS logo tile", "LS logo", "Brand block", "Brand name", "Brand tagline",
            "Slide index", "Title", "Subtitle", "Header rule", "Evidence status", "Labels",
        }
        footer = ("Footer", "Footer rule", "Slide number")
        below = [s for s in slide.shapes if s.name not in chrome and s.name not in footer]
        assert all(s.top >= banner.top + banner.height for s in below)  # content under banner
        assert banner.top + banner.height < prs.slide_height - Inches(0.7)  # above the footer
    assert check_text_bounds(path) == []


def test_short_labels_keep_a_compact_banner(tmp_path, source):
    from pptx.util import Inches

    from launchscout.exports.pptx_builder import build_pptx

    path = tmp_path / "short.pptx"
    build_pptx(deck.apply_labels(_spec(), source), source, path, date(2026, 10, 14))
    first = Presentation(str(path)).slides[0]
    banner = next(s for s in first.shapes if s.name == "Labels")
    assert banner.height < Inches(1.0)


def test_bounds_check_flags_a_banner_too_small_for_its_text(tmp_path, source):
    from launchscout.exports.pptx_builder import build_pptx, check_text_bounds

    path = tmp_path / "small.pptx"
    build_pptx(_banner_deck(source), source, path, date(2026, 10, 14))
    prs = Presentation(str(path))
    for shape in prs.slides[0].shapes:
        if shape.name == "Labels":
            shape.height = 365760  # the old fixed 0.4in strip
    prs.save(str(path))
    assert any("outside its banner" in p for p in check_text_bounds(path))


# ---------------------------------------------------------------- enterprise layout (V2)
def _build(tmp_path, source, spec=None, name="deck.pptx"):
    from launchscout.exports.pptx_builder import build_pptx

    path = tmp_path / name
    build_pptx(deck.apply_labels(spec or _spec(), source), source, path, date(2026, 10, 14))
    return path, Presentation(str(path))


def _more_competitors(source, n):
    template = source.competitors.competitors[1]
    comps = [
        template.model_copy(update={"name": f"Rival {i}", "url": f"https://r{i}.example/"})
        for i in range(n)
    ]
    return source.model_copy(
        update={"competitors": source.competitors.model_copy(update={"competitors": comps})}
    )


def _shape(slide, name):
    return next(s for s in slide.shapes if s.name == name)


def test_every_slide_has_logo_status_footer_and_native_shapes(tmp_path, source):
    path, prs = _build(tmp_path, source)
    assert len(prs.slides) == 6
    for number, slide in enumerate(prs.slides, 1):
        names = [s.name for s in slide.shapes]
        assert names.count("LS logo") == 1 and "LS logo tile" in names
        assert sum(1 for s in slide.shapes if s.shape_type == 13) == 1
        assert _shape(slide, "Slide number").text_frame.text == f"{number:02d} / 06"
        assert _shape(slide, "Brand tagline").text_frame.text == "Competitor Research Crew"
        assert _shape(slide, "Brand name").text_frame.text == "LaunchScout"
        assert "PARTIAL EVIDENCE" in _shape(slide, "Evidence status").text_frame.text
        assert "PARTIAL EVIDENCE" in _shape(slide, "Labels").text_frame.text
        assert slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip()
    assert _shape(prs.slides[0], "LS logo").image.sha1  # a real picture asset


def test_visual_components_present_and_editable(tmp_path, source):
    _, prs = _build(tmp_path, source)
    s1, s2, s3, s4, s5, s6 = prs.slides
    names1 = [s.name for s in s1.shapes]
    assert sum(1 for n in names1 if n.startswith("Journey ") and n.endswith(" icon")) == 4
    charts1 = [s for s in s1.shapes if getattr(s, "has_chart", False) and s.has_chart]
    assert len(charts1) == 2  # native, editable doughnut gauges
    assert sum(1 for s in s2.shapes if s.name.startswith("Competitor: ")) == 2
    assert sum(1 for s in s2.shapes if s.name.startswith("Price: ")) == 2  # status chips
    assert sum(1 for s in s2.shapes if s.name.startswith("Grid ")) >= 4  # evidence footprint
    charts3 = [s for s in s3.shapes if getattr(s, "has_chart", False) and s.has_chart]
    assert len(charts3) == 1 and charts3[0].chart.chart_type is not None  # native theme bars
    assert {"Value proposition", "Positioning tag"} <= {s.name for s in s4.shapes}
    assert sum(1 for s in s4.shapes if s.name.startswith("Flow arrow")) == 3
    assert {"Stage 1", "Stage 2", "Stage 3", "Stage 4"} <= {s.name for s in s4.shapes}
    assert sum(1 for s in s5.shapes if s.name.startswith("Journey stage")) >= 1
    warning = _shape(s6, "Price warning").text_frame.text
    assert warning == "RECOMMENDATION — NOT A VERIFIED MARKET PRICE"
    assert _shape(s6, "Price value").text_frame.text == "₹99"
    for slide in prs.slides:
        assert all(s.shape_type != 13 or s.name == "LS logo" for s in slide.shapes)


def test_layout_has_no_overflow_collisions_or_footer_hits(tmp_path, source):
    from launchscout.exports.pptx_builder import check_text_bounds

    path, _ = _build(tmp_path, source)
    assert check_text_bounds(path) == []


@pytest.mark.parametrize("n", [1, 3, 5])
def test_layout_adapts_to_competitor_count(tmp_path, source, n):
    from launchscout.exports.pptx_builder import check_text_bounds

    path, prs = _build(tmp_path, _more_competitors(source, n), name=f"n{n}.pptx")
    assert len(prs.slides) == 6
    matrix = " ".join(s.name for s in prs.slides[1].shapes)
    assert all(f"Rival {i}" in matrix for i in range(n))
    assert check_text_bounds(path) == []


def test_zero_competitors_and_no_weeks_do_not_crash(tmp_path, source):
    from launchscout.exports.pptx_builder import check_text_bounds

    empty = _more_competitors(source, 0)
    spec = _spec(lambda s: s[5].update(bullets=[]))
    path, prs = _build(tmp_path, empty, spec, "empty.pptx")
    assert len(prs.slides) == 6
    assert check_text_bounds(path) == []


def test_space_utilisation_is_high_on_every_slide(tmp_path, source):
    from launchscout.exports.pptx_builder import layout_utilization

    path, _ = _build(tmp_path, source)
    values = layout_utilization(path)
    assert all(u >= 0.6 for u in values), values


def test_layout_engine_is_not_hardcoded_to_one_idea():
    import launchscout.exports.pptx_builder as builder
    import launchscout.exports.pptx_components as components

    for module in (builder, components):
        code = Path(module.__file__).read_text(encoding="utf-8")
        for needle in ("Pune", "juice", "699", "gym", "iBites", "Earth Organic"):
            assert needle not in code, (module.__name__, needle)


def test_price_parts_split_recommended_price_text():
    from launchscout.exports.pptx_builder import _price_parts

    head, plan, derived, note = _price_parts(
        "₹699 for a seven-day plan of one juice per day (about ₹100 per juice); proposed only"
    )
    assert head == "₹699" and plan == "Seven-day plan of one juice per day"
    assert derived == "about ₹100 per juice" and note == "proposed only"
    assert _price_parts("INR 99 per bottle")[0] == "₹99"
    assert _price_parts("no figure") == ("", "", "", "no figure")


def test_parse_weeks_orders_weeks_and_falls_back_to_steps():
    from launchscout.exports.pptx_builder import parse_weeks

    assert parse_weeks(["Week 1: a", "Week 2 - b"]) == [("Week 1", "a"), ("Week 2", "b")]
    assert parse_weeks(["do x", "do y"]) == [("Step 1", "do x"), ("Step 2", "do y")]


def test_notes_keep_sources_and_full_text_for_the_slide(tmp_path, source):
    _, prs = _build(tmp_path, source)
    notes = prs.slides[1].notes_slide.notes_text_frame.text
    assert "Sources:" in notes and U1 in notes and "Validated slide text:" in notes
    assert "Labels:" in notes


# ---------------------------------------------------------------- badge / chip sizing
def _chips(prs):
    """Single-line status labels: shapes with word wrap disabled and a text label."""
    for number, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                if shape.text_frame.word_wrap is False:
                    yield number, shape


def test_status_badges_never_wrap_and_hold_their_full_label(tmp_path, source):
    _, prs = _build(tmp_path, source)
    seen = {name for _, shape in _chips(prs) for name in [shape.name]}
    assert "Positioning tag" in seen and "Evidence status" in seen
    for number, shape in _chips(prs):
        frame = shape.text_frame
        label = frame.text
        assert "\n" not in label
        size = max(r.font.size.pt for p in frame.paragraphs for r in p.runs)
        inner_pt = shape.width / 12700 - (frame.margin_left + frame.margin_right) / 12700
        # bold capitals need about 0.6 em each; the chip must hold the whole label with margin
        assert inner_pt >= len(label) * size * 0.6, (number, shape.name, label)


def test_badges_fit_a_measured_calibri_width_with_padding(tmp_path, source):
    from PIL import ImageFont

    font_path = Path("C:/Windows/Fonts/calibrib.ttf")
    if not font_path.exists():
        pytest.skip("Calibri is not installed; measured-width check skipped")
    _, prs = _build(tmp_path, source)
    for number, shape in _chips(prs):
        frame = shape.text_frame
        size = max(r.font.size.pt for p in frame.paragraphs for r in p.runs)
        font = ImageFont.truetype(str(font_path), 100)
        measured_pt = font.getlength(frame.text) / 100 * size
        inner_pt = shape.width / 12700 - (frame.margin_left + frame.margin_right) / 12700
        assert inner_pt >= measured_pt * 1.15, (number, shape.name, frame.text)


def test_very_long_badge_label_gets_a_wider_chip_instead_of_wrapping():
    from pptx import Presentation as P
    from pptx.util import Inches

    from launchscout.exports.pptx_components import BURGUNDY, BURGUNDY_LIGHT, chip, chip_width

    prs = P()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    short, long = "RECOMMENDATION", "RECOMMENDATION — PROPOSED, NOT VERIFIED"
    assert chip_width(long, 11) > chip_width(short, 11) * 2
    shape, width = chip(slide, 1, 1, long, BURGUNDY_LIGHT, BURGUNDY, size=11, name="long")
    assert shape.text_frame.word_wrap is False
    assert shape.width == Inches(width) or abs(shape.width - Inches(width)) < 5


# ---------------------------------------------------------------- V3 investor deck
def test_v3_every_validated_bullet_is_preserved_in_the_notes(tmp_path, source):
    spec = deck.apply_labels(_spec(), source)
    _, prs = _build(tmp_path, source)
    for slide, slide_spec in zip(prs.slides, spec.slides, strict=True):
        notes = slide.notes_slide.notes_text_frame.text
        assert all(b in notes for b in slide_spec.bullets)
        assert all(label in notes for label in slide_spec.labels)


def test_v3_unverified_prices_are_labelled_and_recommendations_are_not_prices(tmp_path, source):
    _, prs = _build(tmp_path, source)
    chips = [s.text_frame.text for s in prs.slides[1].shapes if s.name.startswith("Price: ")]
    assert "PRICE UNKNOWN" in chips  # an unverified price is never shown as a figure
    assert len(chips) == 2 and sum(1 for c in chips if c.startswith("₹")) == 1
    s6 = prs.slides[5]
    assert _shape(s6, "Price warning").text_frame.text.startswith("RECOMMENDATION")
    assert "PRICE UNKNOWN" not in " ".join(
        s.text_frame.text for s in s6.shapes if s.has_text_frame
    )


def test_v3_header_footer_and_sources_strip_on_every_slide(tmp_path, source):
    _, prs = _build(tmp_path, source)
    names = (
        "Problem & Opportunity", "Market Landscape", "Customer Insights",
        "Solution & Differentiation", "Go-to-Market Strategy", "Financials & Roadmap",
    )
    for number, (slide, name) in enumerate(zip(prs.slides, names, strict=True), 1):
        assert _shape(slide, "Slide index").text_frame.text == f"{number:02d}."
        assert _shape(slide, "Title").text_frame.text == name
        footer = _shape(slide, "Footer").text_frame.text
        assert footer.startswith("LaunchScout | " + name) and "Partial evidence" in footer
        assert source.run_id not in footer  # a shortened run id is shown; the full id is in notes
        assert source.run_id in slide.notes_slide.notes_text_frame.text
        assert _shape(slide, "Sources").text_frame.text.startswith("Sources:")


def test_v3_compact_rows_for_many_competitors_keep_status_chips(tmp_path, source):
    from launchscout.exports.pptx_builder import check_text_bounds

    path, prs = _build(tmp_path, _more_competitors(source, 7), name="many.pptx")
    assert sum(1 for s in prs.slides[1].shapes if s.name.startswith("Price: ")) == 7
    assert check_text_bounds(path) == []


def test_v4_on_slide_text_is_never_cut_with_an_ellipsis(tmp_path, source):
    _, prs = _build(tmp_path, source)
    cut = [
        (n, s.name)
        for n, slide in enumerate(prs.slides, 1)
        for s in slide.shapes
        if s.has_text_frame and s.text_frame.text.rstrip().endswith("…")
    ]
    assert cut == []


def test_shorten_prefers_a_complete_clause_to_an_ellipsis():
    from launchscout.exports.pptx_components import _shorten

    text = "Confirm feasibility and service areas; decide choice or rotation and validate price"
    assert _shorten(text, 50) == "Confirm feasibility and service areas"
    assert _shorten("a" * 40, 20).endswith("…")  # no clause boundary: the last resort
